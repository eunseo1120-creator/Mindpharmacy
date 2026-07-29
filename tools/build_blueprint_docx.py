from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "full-chapter-blueprint.md"
OUTPUT = ROOT / "docs" / "mind-pharmacy-full-chapter-blueprint.docx"

INK = RGBColor(36, 30, 27)
ACCENT = RGBColor(116, 70, 62)
MUTED = RGBColor(104, 94, 86)
PAPER = "EEE8DD"
FONT = "Arial Unicode MS"


def font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clean(text):
    return text.replace("**", "").replace("`", "")


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.85)
sec.bottom_margin = Inches(0.85)
sec.left_margin = Inches(0.9)
sec.right_margin = Inches(0.9)
sec.header_distance = Inches(0.45)
sec.footer_distance = Inches(0.45)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, before, after in [
    ("Heading 1", 17, 18, 9),
    ("Heading 2", 13.5, 14, 7),
    ("Heading 3", 11.5, 10, 5),
]:
    style = styles[name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = ACCENT
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("MIND PHARMACY  ·  GAME DESIGN"), 8.5, True, MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Prototype v0.2  |  2026.07.29"), 8.5, False, MUTED)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(42)
title.paragraph_format.space_after = Pt(8)
font(title.add_run("마음 약방"), 30, True, INK)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(28)
font(subtitle.add_run("전체 챕터 · 방 구조 · 단서 · 퍼즐 로직 · 스토리 설계안"), 15, False, ACCENT)

lead = doc.add_paragraph()
lead.paragraph_format.left_indent = Inches(0.18)
lead.paragraph_format.right_indent = Inches(0.18)
lead.paragraph_format.space_after = Pt(22)
shade(lead, PAPER)
font(lead.add_run("설계 원칙  "), 10.5, True, ACCENT)
font(lead.add_run("상처를 지우는 탈출이 아니라, 도움을 받으며 다음 행동을 고르는 과정을 플레이하게 한다."), 10.5)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
in_code = False
code_lines = []

for raw in lines[4:]:
    line = raw.rstrip()
    if line.startswith("```"):
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(9)
            shade(p, "F2EFEA")
            font(p.add_run("\n".join(code_lines)), 8.7, False, MUTED)
            code_lines = []
        in_code = not in_code
        continue
    if in_code:
        code_lines.append(line)
        continue
    if not line:
        continue
    if line.startswith("# "):
        continue
    if line.startswith("## "):
        heading = line[3:]
        if heading.startswith(("1장", "2장", "3장", "피날레")):
            doc.add_page_break()
        doc.add_paragraph(heading, style="Heading 1")
    elif line.startswith("### "):
        doc.add_paragraph(line[4:], style="Heading 2")
    elif re.match(r"^\d+\. ", line):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        font(p.add_run(clean(re.sub(r"^\d+\. ", "", line))), 10.5)
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        font(p.add_run(clean(line[2:])), 10.5)
    else:
        p = doc.add_paragraph()
        font(p.add_run(clean(line)), 10.5)

doc.core_properties.title = "마음 약방 전체 챕터 설계안"
doc.core_properties.subject = "방 구조, 시각 요소, 단서, 퍼즐 로직, 스토리"
doc.core_properties.author = "Mind Pharmacy Team"
doc.save(OUTPUT)
print(OUTPUT)
